"""Local document text extraction service (Camino B).

Provides provider-independent text extraction for PDF and Word (.docx) documents.
Extracts pure text from binary streams without relying on external multimodal APIs,
optimizing cost and preparing the foundation for future RAG chunking.
"""

from __future__ import annotations

import io
from pathlib import Path

import docx
import pypdf
import structlog
from fastapi import UploadFile

log = structlog.get_logger()


def extract_text_from_pdf(content: bytes) -> str:
    """Extract plain text from PDF bytes using pypdf.

    Returns extracted text or a fallback marker if unextractable.
    """
    try:
        reader = pypdf.PdfReader(io.BytesIO(content))
        extracted_pages: list[str] = []
        for index, page in enumerate(reader.pages):
            text = page.extract_text() or ""
            if text.strip():
                extracted_pages.append(text.strip())

        if not extracted_pages:
            return "[No extractable text found in PDF]"
        return "\n\n".join(extracted_pages)
    except Exception as exc:
        log.warning("pdf_extraction_failed", error=str(exc))
        return f"[Error extracting PDF content: {exc}]"


def extract_text_from_docx(content: bytes) -> str:
    """Extract text from Word (.docx) bytes using python-docx."""
    try:
        document = docx.Document(io.BytesIO(content))
        extracted_elements: list[str] = []

        for paragraph in document.paragraphs:
            text = paragraph.text.strip()
            if text:
                extracted_elements.append(text)

        for table in document.tables:
            for row in table.rows:
                row_texts = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_texts:
                    extracted_elements.append(" | ".join(row_texts))

        if not extracted_elements:
            return "[No extractable text found in Word document]"
        return "\n".join(extracted_elements)
    except Exception as exc:
        log.warning("docx_extraction_failed", error=str(exc))
        return f"[Error extracting DOCX content: {exc}]"


async def extract_attachment_text(file: UploadFile) -> str:
    """Read and extract text from an uploaded file based on its extension."""
    filename = file.filename or "unknown_file"
    extension = Path(filename).suffix.lower()

    content = await file.read()

    if extension == ".pdf":
        body = extract_text_from_pdf(content)
    elif extension in {".docx", ".doc"}:
        body = extract_text_from_docx(content)
    elif extension in {".txt", ".md", ".json", ".csv"}:
        try:
            body = content.decode("utf-8", errors="replace").strip()
        except Exception as exc:
            log.warning("text_file_decode_failed", error=str(exc), filename=filename)
            body = f"[Error reading text file: {exc}]"
    else:
        body = f"[Unsupported document format '{extension}']"

    return f"--- attachment: {filename} ---\n{body}"
